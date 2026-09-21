"""Every engine walk reads a run's font as the text state holds it.

ISO 32000-2 §9.3.1: the text state holds a font DICTIONARY. `Tf` selects the
one its name gives in the current stream's resources; an ExtGState /Font
entry sets the font and the size with no name at all (Table 57); a form XObject
runs in the state of the Do that draws it, font and rendering mode included,
whatever its own resources call by the same name (§8.10.1); and a pattern cell
starts in the state the stream that owns the pattern started in (§8.7.3.1 b).
Each class below pins one walk that read the font, the size or the rendering
mode another way.
"""

from __future__ import annotations

import io
import os

import pikepdf
import pytest
from pikepdf import Array, Dictionary, Name

from text_state_shapes import TEXT, shape_pdf


def _helvetica(pdf, base: str = "Helvetica"):
    return pdf.make_indirect(Dictionary(
        Type=Name.Font, Subtype=Name.Type1, BaseFont=Name("/" + base),
        Encoding=Name.WinAnsiEncoding))


def _form(pdf, content: bytes, resources=None, bbox=(0, 0, 400, 400)):
    form = pdf.make_stream(content)
    form["/Type"] = Name.XObject
    form["/Subtype"] = Name.Form
    form["/BBox"] = Array(list(bbox))
    form["/Resources"] = resources if resources is not None else Dictionary()
    return pdf.make_indirect(form)


def _save(pdf, tmp_dir: str, name: str) -> str:
    path = os.path.join(tmp_dir, name)
    pdf.save(path)
    pdf.close()
    return path


# ── text extraction (pdfminer's layout analysis) ──────────────────────────


class TestExtraction:
    """pdfminer ignores `gs` and starts every form from an empty text state;
    the engine reads pages through `TextStateInterpreter`."""

    @pytest.mark.parametrize("label", ("A", "B"))
    def test_extract_text_reads_the_run(self, tmp_dir, label):
        from engine.extract_text import extract_text

        text = extract_text(shape_pdf(tmp_dir, label))["text"]
        assert TEXT.decode("ascii") in text

    def test_the_text_export_reads_the_run(self, tmp_dir):
        from engine.text_export import export_text

        out = os.path.join(tmp_dir, "a.txt")
        export_text(shape_pdf(tmp_dir, "A"), out)
        with open(out, encoding="utf-8") as handle:
            assert TEXT.decode("ascii") in handle.read()

    def test_compare_reads_the_run(self, tmp_dir):
        from engine.compare import _extract_lines

        lines, _pages, _count = _extract_lines(shape_pdf(tmp_dir, "A"))
        assert TEXT.decode("ascii") in lines

    def _form_text_page(self, tmp_dir) -> str:
        pdf = pikepdf.new()
        font = _helvetica(pdf)
        page = pdf.add_blank_page(page_size=(400, 400))
        page.Resources = Dictionary(
            Font=Dictionary(F1=font),
            XObject=Dictionary(Fm0=_form(
                pdf, b"BT /F1 12 Tf 60 300 Td (WORDS DRAWN BY A FORM) Tj ET",
                Dictionary(Font=Dictionary(F1=font)))),
        )
        page.Contents = pdf.make_stream(b"BT /F1 12 Tf 60 100 Td (Page words) Tj ET /Fm0 Do")
        return _save(pdf, tmp_dir, "form_text.pdf")

    def test_search_in_files_finds_words_a_form_draws(self, tmp_dir):
        # pdfminer lays a form out as a figure, outside the page's text boxes.
        from engine.search_in_files import search_in_files

        hits = search_in_files([self._form_text_page(tmp_dir)], "drawn by a form")["hits"]
        assert [(hit["page"], hit["count"]) for hit in hits] == [(1, 1)]

    def test_search_in_files_finds_words_drawn_in_an_inherited_font(self, tmp_dir):
        from engine.search_in_files import search_in_files

        hits = search_in_files([shape_pdf(tmp_dir, "B")], "secret words")["hits"]
        assert [(hit["page"], hit["count"]) for hit in hits] == [(1, 1)]

    def test_batch_ocr_counts_the_words_a_form_draws(self, tmp_dir):
        from engine.batch_ocr import _document_page_text

        text = _document_page_text(self._form_text_page(tmp_dir), 1)[0]
        assert "WORDS DRAWN BY A FORM" in text


# ── accessibility ─────────────────────────────────────────────────────────


class TestTheFontsAccessibilityChecks:
    """ISO 14289-1 cl. 7.21 governs the fonts text is drawn WITH."""

    def test_a_font_an_extgstate_sets_is_used(self, tmp_dir):
        from engine.accessibility import _rendered_fonts

        with pikepdf.open(shape_pdf(tmp_dir, "A")) as pdf:
            rendered, unread = _rendered_fonts(pdf)
            names = sorted(str(font.get("/BaseFont")) for font in rendered.values())
        assert unread == []
        assert names == ["/Wide"]

    def test_a_form_draws_with_the_font_it_inherits(self, tmp_dir):
        from engine.accessibility import _rendered_fonts

        with pikepdf.open(shape_pdf(tmp_dir, "B")) as pdf:
            rendered, _unread = _rendered_fonts(pdf)
            names = sorted(str(font.get("/BaseFont")) for font in rendered.values())
        # The form's own /F1 ("Narrow") draws nothing.
        assert names == ["/Wide"]

    def test_a_form_under_the_invisible_mode_uses_no_font(self, tmp_dir):
        from engine.accessibility import _rendered_fonts

        pdf = pikepdf.new()
        font = _helvetica(pdf)
        page = pdf.add_blank_page(page_size=(400, 400))
        page.Resources = Dictionary(
            Font=Dictionary(F1=font),
            XObject=Dictionary(Fm0=_form(
                pdf, b"BT /F1 12 Tf 60 300 Td (recognized words) Tj ET",
                Dictionary(Font=Dictionary(F1=font)))),
        )
        page.Contents = pdf.make_stream(b"3 Tr /Fm0 Do")
        with pikepdf.open(_save(pdf, tmp_dir, "ocr_form.pdf")) as doc:
            rendered, _unread = _rendered_fonts(doc)
        assert rendered == {}

    def test_two_fonts_that_share_a_name_are_two_findings(self, tmp_dir):
        # The page's /F1 and the form's own /F1 are two Type 3 fonts; neither
        # maps its codes to Unicode, and each is a finding of its own.
        from engine.accessibility import _check_character_encoding, _Check, _Pages

        pdf = pikepdf.new()
        page_font = pdf.make_indirect(Dictionary(Type=Name.Font, Subtype=Name.Type3))
        form_font = pdf.make_indirect(Dictionary(Type=Name.Font, Subtype=Name.Type3))
        page = pdf.add_blank_page(page_size=(400, 400))
        page.Resources = Dictionary(
            Font=Dictionary(F1=page_font),
            XObject=Dictionary(Fm0=_form(
                pdf, b"BT /F1 12 Tf 60 300 Td (b) Tj ET",
                Dictionary(Font=Dictionary(F1=form_font)))),
        )
        page.Contents = pdf.make_stream(b"BT /F1 12 Tf 60 100 Td (a) Tj ET /Fm0 Do")
        with pikepdf.open(_save(pdf, tmp_dir, "two_f1.pdf")) as doc:
            check = _Check("character_encoding", "page_content")
            _check_character_encoding(check, _Pages(doc))
        assert len(check.findings) == 2


# ── font inventories ──────────────────────────────────────────────────────


class TestTheFontInventories:
    def test_the_document_font_list_names_an_extgstate_font(self, tmp_dir):
        from engine.font_inventory import list_document_fonts

        names = [font["name"] for font in list_document_fonts(shape_pdf(tmp_dir, "A"))["fonts"]]
        assert "Wide" in names

    def test_the_document_font_list_names_a_soft_mask_s_font(self, tmp_dir):
        from engine.font_inventory import list_document_fonts

        pdf = pikepdf.new()
        group = _form(pdf, b"BT /F1 12 Tf 60 300 Td (mask) Tj ET",
                      Dictionary(Font=Dictionary(F1=_helvetica(pdf, "Courier"))))
        page = pdf.add_blank_page(page_size=(400, 400))
        page.Resources = Dictionary(ExtGState=Dictionary(GM=Dictionary(
            Type=Name.ExtGState,
            SMask=Dictionary(Type=Name.Mask, S=Name.Luminosity, G=group))))
        page.Contents = pdf.make_stream(b"/GM gs 0 0 100 100 re f")
        names = [font["name"] for font in list_document_fonts(_save(pdf, tmp_dir, "mask.pdf"))["fonts"]]
        assert names == ["Courier"]

    def test_preflight_s_resource_walk_reaches_an_extgstate_font(self, tmp_dir):
        from engine.preflight import walk_page_resources

        reached = []
        with pikepdf.open(shape_pdf(tmp_dir, "A")) as pdf:
            walk_page_resources(
                pdf.pages[0], on_font=lambda font, category: reached.append(str(font.get("/BaseFont")))
            )
        assert "/Wide" in reached


# ── slide export ──────────────────────────────────────────────────────────


class TestTheSlideFont:
    """A text box takes its face from the font the text was drawn in, not
    from whatever the page's resources call by the same name."""

    @staticmethod
    def _faces(path: str, tmp_dir: str, gs_path: str) -> dict:
        from pptx import Presentation

        from engine.slide_export import export_slides

        out = os.path.join(tmp_dir, "deck.pptx")
        export_slides(path, out, gs_path=gs_path)
        faces = {}
        for shape in Presentation(out).slides[0].shapes:
            if shape.has_text_frame and shape.text_frame.text:
                run = shape.text_frame.paragraphs[0].runs[0]
                faces[shape.text_frame.text] = (run.font.name, bool(run.font.bold))
        return faces

    def test_a_form_s_own_font_names_the_face(self, tmp_dir, gs_path):
        pdf = pikepdf.new()
        page = pdf.add_blank_page(page_size=(400, 400))
        page.Resources = Dictionary(
            Font=Dictionary(F1=_helvetica(pdf)),
            XObject=Dictionary(Fm0=_form(
                pdf, b"BT /F1 12 Tf 60 300 Td (Serif words) Tj ET",
                Dictionary(Font=Dictionary(F1=_helvetica(pdf, "Times-Bold"))))),
        )
        page.Contents = pdf.make_stream(b"BT /F1 12 Tf 60 100 Td (Sans words) Tj ET /Fm0 Do")
        faces = self._faces(_save(pdf, tmp_dir, "faces.pdf"), tmp_dir, gs_path)
        assert faces["Sans words"] == ("Arial", False)
        assert faces["Serif words"] == ("Times New Roman", True)

    def test_an_extgstate_font_names_the_face(self, tmp_dir, gs_path):
        faces = self._faces(shape_pdf(tmp_dir, "A"), tmp_dir, gs_path)
        assert faces[TEXT.decode("ascii")] == ("Wide", False)

    def test_a_face_name_that_is_not_utf8_exports_under_the_fallback_face(self, tmp_dir, gs_path):
        # A name reads as text only through UTF-8 (ISO 32000-2 §7.3.5), so a
        # /BaseFont with a byte that is not UTF-8 names no installed face. A
        # strict decode of such a name raises and stops the export.
        pdf = pikepdf.new()
        page = pdf.add_blank_page(page_size=(400, 400))
        odd = pdf.make_indirect(Dictionary(
            Type=Name.Font, Subtype=Name.Type1,
            BaseFont=pikepdf.Object.parse(b"/ABCDEF+FreeMonoBold#c4"),
            Encoding=Name.WinAnsiEncoding, FirstChar=32, LastChar=126, Widths=[600] * 95))
        page.Resources = Dictionary(Font=Dictionary(F1=_helvetica(pdf), F2=odd))
        page.Contents = pdf.make_stream(
            b"BT /F1 12 Tf 60 300 Td (Plain words) Tj ET BT /F2 12 Tf 60 200 Td (Odd words) Tj ET")
        faces = self._faces(_save(pdf, tmp_dir, "odd-name.pdf"), tmp_dir, gs_path)
        assert faces.pop("Plain words") == ("Arial", False)
        assert list(faces.values()) == [("Arial", True)]


# ── scanned-page compression ──────────────────────────────────────────────


class TestTheScanCheck:
    def test_text_a_form_draws_under_the_invisible_mode_is_invisible(self, tmp_dir):
        from engine.mrc import _has_other_visible_content
        from engine.redact import _resolve_resources

        pdf = pikepdf.new()
        font = _helvetica(pdf)
        page = pdf.add_blank_page(page_size=(400, 400))
        page.Resources = Dictionary(
            Font=Dictionary(F1=font),
            XObject=Dictionary(Fm0=_form(
                pdf, b"BT /F1 12 Tf 60 300 Td (recognized words) Tj ET",
                Dictionary(Font=Dictionary(F1=font)))),
        )
        page.Contents = pdf.make_stream(b"q 3 Tr /Fm0 Do Q")
        with pikepdf.open(_save(pdf, tmp_dir, "ocr.pdf")) as doc:
            target = doc.pages[0]
            visible = _has_other_visible_content(
                doc, list(pikepdf.parse_content_stream(target)), _resolve_resources(target), None
            )
        assert visible is False


# ── automatic tagging ─────────────────────────────────────────────────────


class TestAutotagSizes:
    """A block's size is the size the text state holds at its shows, which a
    `Tf` before the block or an ExtGState sets as surely as one inside it."""

    def _tally(self, tmp_dir, content: bytes, name: str) -> dict:
        from engine.autotag import autotag

        pdf = pikepdf.new()
        page = pdf.add_blank_page(page_size=(612, 792))
        page.Resources = Dictionary(
            Font=Dictionary(F1=_helvetica(pdf)),
            ExtGState=Dictionary(GH=Dictionary(
                Type=Name.ExtGState, Font=Array([_helvetica(pdf, "Helvetica-Bold"), 24]))),
        )
        page.Contents = pdf.make_stream(content)
        src = _save(pdf, tmp_dir, name)
        return autotag(src, os.path.join(tmp_dir, "tagged-" + name))

    def test_a_tf_before_the_block_sizes_it(self, tmp_dir):
        tally = self._tally(
            tmp_dir,
            b"/F1 24 Tf BT 72 700 Td (A heading) Tj ET\n"
            b"/F1 12 Tf BT 72 650 Td (Body text that runs long enough to vote) Tj ET\n"
            b"BT 72 630 Td (More body text that runs long enough to vote) Tj ET\n",
            "before_bt.pdf",
        )
        assert tally["headings"] == 1
        assert tally["paragraphs"] == 2

    def test_an_extgstate_font_sizes_the_block(self, tmp_dir):
        tally = self._tally(
            tmp_dir,
            b"BT /GH gs 72 700 Td (A heading) Tj ET\n"
            b"BT /F1 12 Tf 72 650 Td (Body text that runs long enough to vote) Tj ET\n"
            b"BT 72 630 Td (More body text that runs long enough to vote) Tj ET\n",
            "gs.pdf",
        )
        assert tally["headings"] == 1
        assert tally["paragraphs"] == 2


# ── redaction and pattern cells ───────────────────────────────────────────


def _pattern_page(pdf, font, cell_content: bytes):
    """A page whose form fills its box with a one-tile pattern. The cell has
    no Tf of its own: it draws in the font the form started in, the page's
    /F1, selected before the Do."""
    cell = pdf.make_stream(cell_content)
    cell["/Type"] = Name.Pattern
    cell["/PatternType"] = 1
    cell["/PaintType"] = 1
    cell["/TilingType"] = 1
    cell["/BBox"] = Array([0, 0, 400, 400])
    cell["/XStep"] = 400
    cell["/YStep"] = 400
    cell["/Resources"] = Dictionary()
    form = _form(pdf, b"/Pattern cs /P0 scn 0 0 400 400 re f",
                 Dictionary(Pattern=Dictionary(P0=pdf.make_indirect(cell))))
    page = pdf.add_blank_page(page_size=(400, 400))
    page.Resources = Dictionary(Font=Dictionary(F1=font), XObject=Dictionary(Fm0=form))
    return page


class TestAPatternCellDrawsInTheStateItsOwnerStartedIn:
    def test_a_mark_splits_the_cell_s_run_at_the_inherited_font_s_advances(self, tmp_dir):
        from test_redact_text_state import ADVANCE, MARK, _shows, _simple_font

        from engine.redact import redact

        pdf = pikepdf.new()
        page = _pattern_page(pdf, _simple_font(pdf, ADVANCE, "Wide"),
                             b"BT 60 300 Td (" + TEXT + b") Tj ET")
        page.Contents = pdf.make_stream(b"BT /F1 12 Tf ET /Fm0 Do")
        src = _save(pdf, tmp_dir, "cell.pdf")
        out = os.path.join(tmp_dir, "cell-out.pdf")
        redact(src, out, [{"page": 1, "rect": MARK}])
        with pikepdf.open(out) as doc:
            form = next(iter(doc.pages[0].Resources.XObject.values()))
            cell = next(iter(form.Resources.Pattern.values()))
            shows = [
                [bytes(p) if isinstance(p, pikepdf.String) else float(p) for p in ins.operands[0]]
                if str(ins.operator) == "TJ" else [bytes(ins.operands[-1])]
                for ins in pikepdf.parse_content_stream(cell)
                if str(ins.operator) in ("Tj", "TJ")
            ]
        assert shows == [[b"PUBLIC ", -6.0 * ADVANCE, b" WORDS"]]
        assert _shows(out) == []

    def test_the_font_cut_after_a_redaction_keeps_the_cell_s_glyphs(self, tmp_dir):
        # A redaction cuts a font to the glyphs still drawn. The cell's
        # "Wave" draws in the page's /F1, so its glyphs stay.
        from fontTools.ttLib import TTFont

        from engine.redact import redact
        from outline_builders import embed_truetype, fonts_available

        if not fonts_available():
            pytest.skip("bundled fonts not provisioned")
        pdf = pikepdf.new()
        page = _pattern_page(pdf, embed_truetype(pdf), b"BT 1 0 0 1 10 40 Tm (Wave) Tj ET")
        page.Contents = pdf.make_stream(
            b"BT /F1 24 Tf 1 0 0 1 20 300 Tm (SECRET) Tj ET BT /F1 24 Tf ET /Fm0 Do")
        src = _save(pdf, tmp_dir, "cut.pdf")
        out = os.path.join(tmp_dir, "cut-out.pdf")
        redact(src, out, [{"page": 1, "rect": [15, 290, 200, 330]}])
        with pikepdf.open(out) as doc:
            program = doc.pages[0].Resources.Font.F1.FontDescriptor.FontFile2.read_bytes()
        face = TTFont(io.BytesIO(program))
        cmap, glyphs = face.getBestCmap(), face["glyf"]
        for ch in "Wave":
            assert glyphs[cmap[ord(ch)]].numberOfContours > 0, ch


# ── the paragraph layer ───────────────────────────────────────────────────


class TestTheParagraphLayerReadsTheDrawnFont:
    def test_an_edit_two_forms_deep_keeps_the_font_s_kerning(self, tmp_dir):
        # The inner form's own resources lack /F9; the name resolves in the
        # outer form's, the stream that invokes it. The page defines no /F9.
        # The edit's layout reads the kerning of the font the run was drawn
        # in, so A|V stays kerned.
        from outline_builders import FONT_DIR, embed_truetype, fonts_available

        from engine.text_paragraphs import list_text_paragraphs, replace_paragraph_text

        if not fonts_available():
            pytest.skip("bundled fonts not provisioned")
        pdf = pikepdf.new()
        serif = embed_truetype(pdf, os.path.join(FONT_DIR, "LiberationSerif-Regular.ttf"), "/LibSerif")
        inner = _form(pdf, b"BT /F9 24 Tf 72 700 Td (AVAV WAVE) Tj ET", Dictionary(),
                      bbox=(0, 0, 612, 792))
        outer = _form(pdf, b"/Fm1 Do", Dictionary(Font=Dictionary(F9=serif), XObject=Dictionary(Fm1=inner)),
                      bbox=(0, 0, 612, 792))
        page = pdf.add_blank_page(page_size=(612, 792))
        page.Resources = Dictionary(XObject=Dictionary(Fm0=outer))
        page.Contents = pdf.make_stream(b"/Fm0 Do")
        src = _save(pdf, tmp_dir, "nested.pdf")
        (para,) = list_text_paragraphs(src, 1)["paragraphs"]
        assert para["editable"] is True
        out = os.path.join(tmp_dir, "nested-out.pdf")
        text = "AVAV WAVES"
        replace_paragraph_text(src, out, 1, para["index"], text,
                               [{"start": 0, "end": len(text), "run": para["runs"][0]}],
                               para["runs"], para["text"], font_path=FONT_DIR)
        with pikepdf.open(out) as doc:
            outer_copy = next(iter(doc.pages[0].Resources.XObject.values()))
            inner_copy = next(iter(outer_copy.Resources.XObject.values()))
            arrays = [ins.operands[0] for ins in pikepdf.parse_content_stream(inner_copy)
                      if str(ins.operator) == "TJ"]
        kerns = [float(v) for array in arrays for v in array if not isinstance(v, pikepdf.String)]
        # Liberation Serif kerns A|V by -264/2048 em: +128.90625 in a TJ.
        assert 128.90625 in [round(v, 5) for v in kerns]
